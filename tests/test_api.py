import sys
import io
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient
from app.api.main import app
import cv2
import numpy as np
import docx
from reportlab.pdfgen import canvas

client = TestClient(app)


def get_test_image_bytes():
    img = np.ones((400, 600, 3), dtype=np.uint8) * 255
    cv2.putText(img, "TEST OCR INVOICE", (50, 80), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 0, 0), 2)
    cv2.putText(img, "Invoice: INV-2026-001", (50, 130), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 0, 0), 1)
    cv2.putText(img, "Total: $500.00", (50, 180), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 0, 0), 1)
    _, encoded = cv2.imencode(".png", img)
    return encoded.tobytes()


def get_test_pdf_bytes():
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.setFont("Helvetica", 14)
    c.drawString(100, 750, "Quarterly Financial Overview Report for Q3 2026")
    c.drawString(100, 720, "Total Gross Revenue: $4,500,000.00")
    c.drawString(100, 690, "Invoice Number: INV-9921-2026")
    c.save()
    buf.seek(0)
    return buf.read()


def get_test_docx_bytes():
    doc = docx.Document()
    doc.add_heading("Corporate Strategy Document", 0)
    doc.add_paragraph("This is the official enterprise policy document for 2026.")
    doc.add_paragraph("Invoice Ref: INV-DOCX-7712 with Total Amount: $1,250.00")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Department"
    t.cell(0, 1).text = "Budget"
    t.cell(1, 0).text = "Engineering"
    t.cell(1, 1).text = "$80,000"
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def test_health_endpoint():
    response = client.get("/api/v1/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    assert "version" in data
    assert "device" in data
    assert data["device"] == "cpu"


def test_ocr_image_endpoint():
    img_bytes = get_test_image_bytes()
    response = client.post(
        "/api/v1/ocr/image",
        files={"file": ("test.png", img_bytes, "image/png")},
        data={"lang": "en", "preprocess": "true"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "results" in data
    assert "full_text" in data


def test_document_upload_image():
    img_bytes = get_test_image_bytes()
    response = client.post(
        "/api/v1/ocr/upload",
        files={"file": ("invoice.png", img_bytes, "image/png")},
        data={"lang": "en"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["file_type"] == "image"
    assert len(data["full_text"]) > 0


def test_document_upload_pdf():
    pdf_bytes = get_test_pdf_bytes()
    response = client.post(
        "/api/v1/ocr/upload",
        files={"file": ("report.pdf", pdf_bytes, "application/pdf")},
        data={"lang": "en"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["file_type"] == "pdf"
    assert data["total_pages"] == 1
    assert "Quarterly Financial" in data["full_text"]


def test_document_upload_docx():
    docx_bytes = get_test_docx_bytes()
    response = client.post(
        "/api/v1/ocr/upload",
        files={"file": ("strategy.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"lang": "en"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert data["file_type"] == "docx"
    assert "Corporate Strategy" in data["full_text"]


def test_ocr_structured_endpoint():
    img_bytes = get_test_image_bytes()
    response = client.post(
        "/api/v1/ocr/structured",
        files={"file": ("test.png", img_bytes, "image/png")},
        data={"lang": "en"}
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "entities" in data


if __name__ == "__main__":
    test_health_endpoint()
    test_ocr_image_endpoint()
    test_document_upload_image()
    test_document_upload_pdf()
    test_document_upload_docx()
    test_ocr_structured_endpoint()
    print("All Comprehensive API Integration Tests Passed!")
